from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE


# =========================================================
# OT Practice Copilot
# PowerPoint生成機能
# =========================================================


# =========================================================
# 基本設定
# =========================================================

SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)


PRESENTATION_TYPES = {
    "学会発表": {
        "subtitle": "Academic Presentation",
    },
    "研修会・セミナー": {
        "subtitle": "Seminar / Workshop",
    },
    "大学講義": {
        "subtitle": "Lecture",
    },
}


# =========================================================
# プレゼンテーション作成
# =========================================================

def create_presentation(
    slides,
    output_path,
    presentation_title="OT Practice Copilot",
    presentation_type="研修会・セミナー",
    subtitle=None,
    author=None,
):
    """
    slides:
        以下のような辞書のリスト

        [
            {
                "title": "スライドタイトル",
                "bullets": [
                    "ポイント1",
                    "ポイント2",
                ],
                "speaker_notes": "講師用ノート",
            }
        ]

    output_path:
        保存するpptxファイルのパス
    """

    output_path = Path(output_path)

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    prs = Presentation()

    # 16:9
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT


    # =====================================================
    # タイトルスライド
    # =====================================================

    title_layout = prs.slide_layouts[0]

    title_slide = prs.slides.add_slide(
        title_layout
    )

    title_slide.shapes.title.text = (
        presentation_title
    )

    if subtitle is None:

        type_info = PRESENTATION_TYPES.get(
            presentation_type,
            {},
        )

        subtitle = type_info.get(
            "subtitle",
            "",
        )

    subtitle_parts = []

    if subtitle:
        subtitle_parts.append(
            subtitle
        )

    if author:
        subtitle_parts.append(
            author
        )

    if (
        len(title_slide.placeholders)
        > 1
    ):

        title_slide.placeholders[
            1
        ].text = "\n".join(
            subtitle_parts
        )


    # =====================================================
    # 本文スライド
    # =====================================================

    for index, slide_data in enumerate(
        slides,
        start=1,
    ):

        add_content_slide(
            prs=prs,
            title=slide_data.get(
                "title",
                f"Slide {index}",
            ),
            bullets=slide_data.get(
                "bullets",
                [],
            ),
            speaker_notes=slide_data.get(
                "speaker_notes",
                "",
            ),
            slide_number=index,
        )


    # =====================================================
    # 保存
    # =====================================================

    prs.save(
        str(output_path)
    )

    return output_path


# =========================================================
# 本文スライド
# =========================================================

def add_content_slide(
    prs,
    title,
    bullets,
    speaker_notes="",
    slide_number=None,
):
    """
    タイトル＋本文の標準スライドを追加する。
    """

    layout = prs.slide_layouts[1]

    slide = prs.slides.add_slide(
        layout
    )


    # -----------------------------------------------------
    # タイトル
    # -----------------------------------------------------

    title_shape = slide.shapes.title

    title_shape.text = title

    if title_shape.text_frame.paragraphs:

        title_paragraph = (
            title_shape
            .text_frame
            .paragraphs[0]
        )

        title_paragraph.font.size = (
            Pt(28)
        )


    # -----------------------------------------------------
    # 本文
    # -----------------------------------------------------

    body = slide.placeholders[1]

    text_frame = body.text_frame

    text_frame.clear()

    for index, bullet in enumerate(
        bullets
    ):

        if index == 0:

            paragraph = (
                text_frame.paragraphs[0]
            )

        else:

            paragraph = (
                text_frame.add_paragraph()
            )

        # 入れ子にも対応
        if isinstance(
            bullet,
            dict,
        ):

            paragraph.text = str(
                bullet.get(
                    "text",
                    "",
                )
            )

            paragraph.level = int(
                bullet.get(
                    "level",
                    0,
                )
            )

        else:

            paragraph.text = str(
                bullet
            )

            paragraph.level = 0

        paragraph.font.size = (
            Pt(22)
        )

        paragraph.space_after = (
            Pt(8)
        )


    # -----------------------------------------------------
    # スライド番号
    # -----------------------------------------------------

    if slide_number is not None:

        add_slide_number(
            slide,
            slide_number,
        )


    # -----------------------------------------------------
    # 講師・発表者ノート
    # -----------------------------------------------------

    if speaker_notes:

        add_speaker_notes(
            slide,
            speaker_notes,
        )

    return slide


# =========================================================
# セクションスライド
# =========================================================

def add_section_slide(
    prs,
    title,
    subtitle="",
):
    """
    章の切り替えに使用するスライド。
    """

    layout = prs.slide_layouts[5]

    slide = prs.slides.add_slide(
        layout
    )

    title_box = slide.shapes.add_textbox(
        Inches(1.0),
        Inches(2.3),
        Inches(11.3),
        Inches(1.0),
    )

    title_frame = (
        title_box.text_frame
    )

    title_frame.text = title

    title_paragraph = (
        title_frame.paragraphs[0]
    )

    title_paragraph.font.size = (
        Pt(32)
    )

    title_paragraph.font.bold = True

    title_paragraph.alignment = (
        PP_ALIGN.CENTER
    )


    if subtitle:

        subtitle_box = (
            slide.shapes.add_textbox(
                Inches(1.5),
                Inches(3.4),
                Inches(10.3),
                Inches(0.7),
            )
        )

        subtitle_frame = (
            subtitle_box.text_frame
        )

        subtitle_frame.text = subtitle

        subtitle_paragraph = (
            subtitle_frame.paragraphs[0]
        )

        subtitle_paragraph.font.size = (
            Pt(20)
        )

        subtitle_paragraph.alignment = (
            PP_ALIGN.CENTER
        )

    return slide


# =========================================================
# スライド番号
# =========================================================

def add_slide_number(
    slide,
    slide_number,
):
    """
    右下に小さくスライド番号を表示する。
    """

    textbox = slide.shapes.add_textbox(
        Inches(12.2),
        Inches(7.0),
        Inches(0.6),
        Inches(0.3),
    )

    text_frame = textbox.text_frame

    text_frame.text = str(
        slide_number
    )

    paragraph = (
        text_frame.paragraphs[0]
    )

    paragraph.font.size = (
        Pt(10)
    )

    paragraph.alignment = (
        PP_ALIGN.RIGHT
    )


# =========================================================
# Speaker Notes
# =========================================================

def add_speaker_notes(
    slide,
    notes_text,
):
    """
    PowerPointのノート欄に
    発表者・講師用ノートを書き込む。
    """

    try:

        notes_slide = (
            slide.notes_slide
        )

        notes_text_frame = (
            notes_slide.notes_text_frame
        )

        if notes_text_frame is not None:

            notes_text_frame.text = (
                notes_text
            )

    except Exception:

        # Notesが利用できない環境でも
        # PPTX生成自体は止めない
        pass


# =========================================================
# 最終まとめスライド
# =========================================================

def add_summary_slide(
    prs,
    title="まとめ",
    points=None,
):
    """
    まとめ用スライド。
    """

    if points is None:
        points = []

    return add_content_slide(
        prs=prs,
        title=title,
        bullets=points,
        speaker_notes="",
    )


# =========================================================
# 注意書きスライド
# =========================================================

def add_notice_slide(
    prs,
    title,
    message,
):
    """
    注意事項や免責事項を表示する。
    """

    layout = prs.slide_layouts[5]

    slide = prs.slides.add_slide(
        layout
    )

    box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(1.0),
        Inches(2.0),
        Inches(11.3),
        Inches(3.0),
    )

    text_frame = box.text_frame

    text_frame.clear()

    paragraph = (
        text_frame.paragraphs[0]
    )

    paragraph.text = title
    paragraph.font.size = Pt(26)
    paragraph.font.bold = True
    paragraph.alignment = (
        PP_ALIGN.CENTER
    )

    paragraph = (
        text_frame.add_paragraph()
    )

    paragraph.text = message
    paragraph.font.size = Pt(18)
    paragraph.alignment = (
        PP_ALIGN.CENTER
    )

    return slide